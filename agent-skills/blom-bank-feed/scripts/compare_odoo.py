#!/usr/bin/env python3
"""
compare_odoo.py — reconcile what the bank says against what Odoo holds.

This is the safety net. Before adding, deleting or re-grouping anything on a bank
journal, run this: it tells you exactly which transactions are missing from Odoo,
which are duplicated, and which are merely mis-dated. It never touches Odoo.

Usage:
    python3 compare_odoo.py --bank jan.xlsx feb.xlsx mar.xlsx \
                            --odoo odoo_lines.csv \
                            --out comparison.xlsx

  --bank   one or more .xlsx files produced by blom2odoo.py
  --odoo   a CSV of the Odoo statement lines, columns: id,date,amount,payment_ref

To produce the Odoo CSV, query account.bank.statement.line for the journal and
date range (fields: date, amount, payment_ref, id) and save the result. See
references/odoo_ids.md for the exact probe call and the trick for reading large
results back out of the Make data store.

Matching is on NORMALISED DESCRIPTION + AMOUNT, deliberately ignoring the date,
because ad-hoc imports frequently map the bank's *value* date instead of its
*business* date. Dates are then reported separately as mismatches rather than
being allowed to break the match. Multiplicity is respected: if the bank shows
three identical -1,000 withdrawals on one day and Odoo holds four, exactly one is
reported as extra.
"""

import argparse, csv, re, sys
from collections import defaultdict, Counter
from decimal import Decimal


def norm(s):
    """Normalise a description for comparison. Strips the date prefix that some
    earlier imports prepended to the label, e.g. '30/01/2026 POS Monthly Fees'."""
    s = str(s or "").upper().strip()
    s = re.sub(r"^\d{2}/\d{2}/\d{4}\s+", "", s)
    return re.sub(r"\s+", " ", s)


def load_bank(paths):
    from openpyxl import load_workbook
    out = []
    for p in paths:
        ws = load_workbook(p).active
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[1] is None:
                continue
            out.append({"src": p, "date": str(r[1]), "label": norm(r[2]),
                        "raw": r[2], "amt": Decimal(str(r[3])),
                        "ref": r[4], "ext": r[0]})
    return out


def load_odoo(path):
    out = []
    with open(path, newline="") as fh:
        for r in csv.DictReader(fh):
            out.append({"id": r["id"], "date": r["date"], "label": norm(r["payment_ref"]),
                        "raw": r["payment_ref"], "amt": Decimal(r["amount"])})
    return out


def compare(bank, odoo):
    pool = defaultdict(list)
    for o in odoo:
        pool[(o["label"], o["amt"])].append(o)
    matched, missing = [], []
    for b in bank:
        k = (b["label"], b["amt"])
        if pool.get(k):
            matched.append((b, pool[k].pop(0)))
        else:
            missing.append(b)
    extra = [o for v in pool.values() for o in v]
    datemis = [(b, o) for b, o in matched if b["date"] != o["date"]]
    return matched, missing, extra, datemis


def write_xlsx(path, bank, odoo, matched, missing, extra, datemis):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    F = "Arial"
    hf = Font(name=F, bold=True, size=10, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F3864")
    base = Font(name=F, size=10); tit = Font(name=F, bold=True, size=14)
    note = Font(name=F, size=9, italic=True, color="666666")
    tot = Font(name=F, bold=True, size=10)
    thin = Border(bottom=Side(style="thin", color="D9D9D9"))

    def head(ws, cols, row=1):
        for i, h in enumerate(cols, 1):
            c = ws.cell(row, i, h); c.font = hf; c.fill = hfill
            c.alignment = Alignment(horizontal="center")

    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    ws["A1"] = "Bank statement vs Odoo"; ws["A1"].font = tit
    ws["A2"] = ("Matched on description + amount. Date deliberately ignored when matching, "
                "then reported separately — ad-hoc imports often map value date instead of "
                "business date.")
    ws["A2"].font = note
    head(ws, ["Measure", "Count", "USD"], 4)
    data = [("Bank transactions", len(bank), sum(x["amt"] for x in bank)),
            ("Lines in Odoo", len(odoo), sum(x["amt"] for x in odoo)),
            ("Matched", len(matched), sum(b["amt"] for b, _ in matched)),
            ("MISSING from Odoo", len(missing), sum(x["amt"] for x in missing)),
            ("EXTRA in Odoo", len(extra), sum(x["amt"] for x in extra)),
            ("Matched, wrong date", len(datemis), None)]
    for i, (k, n, v) in enumerate(data, 5):
        ws.cell(i, 1, k).font = tot if k.isupper() or "MISS" in k or "EXTRA" in k else base
        ws.cell(i, 2, n).font = base
        if v is not None:
            c = ws.cell(i, 3, float(v)); c.font = base; c.number_format = "#,##0.00"
    for col, w in zip("ABC", (40, 12, 16)):
        ws.column_dimensions[col].width = w

    def sheet(name, cols, rowsrc):
        ws = wb.create_sheet(name); head(ws, cols)
        for j, vals in enumerate(rowsrc, 2):
            for i, v in enumerate(vals, 1):
                c = ws.cell(j, i, v); c.font = base; c.border = thin
        ws.freeze_panes = "A2"
        if len(rowsrc):
            ws.auto_filter.ref = f"A1:{chr(64+len(cols))}{len(rowsrc)+1}"
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[chr(64 + i)].width = 30 if i in (3, 4) else 14

    sheet("Missing from Odoo", ["Date", "Description", "Amount", "Bank ref", "External ID"],
          [(b["date"], b["raw"], float(b["amt"]), b["ref"], b["ext"]) for b in missing])
    sheet("Extra in Odoo", ["Odoo line ID", "Date", "Description", "Amount"],
          [(int(o["id"]), o["date"], o["raw"], float(o["amt"])) for o in extra])
    sheet("Date mismatches", ["Odoo line ID", "Odoo date", "Bank date", "Description", "Amount"],
          [(int(o["id"]), o["date"], b["date"], b["raw"], float(b["amt"])) for b, o in datemis])
    sheet("Matched", ["Odoo line ID", "Bank date", "Odoo date", "Description", "Amount"],
          [(int(o["id"]), b["date"], o["date"], b["raw"], float(b["amt"])) for b, o in matched])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bank", nargs="+", required=True)
    ap.add_argument("--odoo", required=True)
    ap.add_argument("--out", default="comparison.xlsx")
    args = ap.parse_args()

    bank = load_bank(args.bank); odoo = load_odoo(args.odoo)
    matched, missing, extra, datemis = compare(bank, odoo)

    print(f"bank {len(bank)} rows, net {sum(x['amt'] for x in bank)}")
    print(f"odoo {len(odoo)} rows, net {sum(x['amt'] for x in odoo)}")
    print(f"\n  matched            {len(matched)}")
    print(f"  MISSING from Odoo  {len(missing)}   {sum(x['amt'] for x in missing)}")
    print(f"  EXTRA in Odoo      {len(extra)}   {sum(x['amt'] for x in extra)}")
    print(f"  wrong date         {len(datemis)}")
    if missing:
        print("\n  most common missing:")
        for lbl, c in Counter(x["label"] for x in missing).most_common(10):
            print(f"     {c:>3}x  {lbl[:56]}")
    if extra:
        print("\n  extra in Odoo (candidates for deletion):")
        for o in extra[:20]:
            print(f"     id {o['id']:>6}  {o['date']}  {o['amt']:>12}  {o['raw'][:44]}")

    write_xlsx(args.out, bank, odoo, matched, missing, extra, datemis)
    print(f"\nwrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
