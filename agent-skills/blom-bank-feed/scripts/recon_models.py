#!/usr/bin/env python3
"""
Reconciliation models for the BLOM USD journal (23) — Presentail SAL, company 2.

Accounts are VERIFIED Odoo internal ids with their company-2 code confirmed via
account.code.mapping. Database is Odoo 19.0, so the field is `trigger`
('manual' | 'auto_reconcile'), NOT the Odoo 18 `rule_type` / `auto_reconcile`.

    python3 recon_models.py apr.xlsx may.xlsx jun.xlsx
"""
import re, sys
from collections import Counter

ACCOUNTS = {
    515:  ("673900",  "Bank charges & commissions",             "expense"),
    513:  ("673601",  "Interest - banks & financial operation",  "expense"),
    1905: ("5300.4",  "Cash USD",                                "asset_cash"),
    1918: ("5121.3",  "POS Blom USD",                            "asset_cash"),
    1559: ("4211",    "Salaries and Wages due to Personnel",     "asset_receivable"),
    516:  ("675100",  "Difference on exchange - negative",       "expense"),
    1661: ("6111",    "Purchase of Raw Materials",               "expense"),
    1903: ("5121.5",  "Blom Bank USD - Business Plus",           "asset_cash"),
}

EXISTING_ON_J23 = {
    20: ("Purchase of Raw Materials", r"TOTERS\ USD\d+\.\d+", 1661),
    21: ("Cash USD",                  r"ATM\ WITHDRAWAL\ ZBLMN\d+", 1905),
    22: ("Transportation",            r"TOTERS\ USD\d+\.\d+", 1885),
    42: ("Loan Zakaria Ammache",      r"IPO/\d+", 1927),
    47: ("Salaries and Wages",        r"SALARY", 1703),
}

# seq, name, regex, counterpart account id (None = no write-off), trigger, note
MODELS = [
    (10, "BLOM: bank fees & commissions",
     r"^(ACCOUNT CHARGES|STATEMENT FEES|STAMPS CHARGES|HIGHEST DB|POS MONTHLY FEES|SWIFT CHARGES|OUTGOING TRANSFER COMMISSION|COMMISSION ON CHECK WITHDRAWAL)",
     515, "auto_reconcile",
     "Existing model 6 'Bank Fees' posts to 383 (601101) and is not journal-restricted; "
     "retire or rescope it so it cannot compete."),

    (20, "BLOM: incoming-transfer commission", r"^COMMISSION\s*/\s*INCOMING", 515, "auto_reconcile",
     "The \\s* is the point: BLOM writes both 'Commission/Incoming Tfr' and "
     "'Commission /Incoming Transfer'. The retired Make rule matched only the first."),

    (30, "BLOM: ATM cash fees", r"^ATM CASH WITHDRAWAL FEE|^ATM CASH DEP\S*\s.*COMM", 515, "auto_reconcile",
     "Must sit ABOVE seq 60/70. The deposit arm is deliberately narrow ('...Comm.') — a loose "
     "'^ATM CASH DEP' misroutes the cash deposit into fees."),

    (40, "BLOM: Alfa / MIC line commission", r"^MIC \d+ \(ALFA\) COMMISSION", 515, "auto_reconcile", ""),

    (45, "BLOM: audit confirmation fee", r"^COMM AUDIT CONF", 515, "auto_reconcile",
     "Annual, around the audit. April 2026: -50.00. The old Make whitelist dropped it."),

    (50, "BLOM: debit interest", r"^DEBIT INTEREST", 513, "auto_reconcile",
     "Existing model 52 already does this to 513 but is not scoped to journal 23. Rescope 52 "
     "rather than creating a duplicate."),

    (60, "BLOM: ATM cash withdrawal", r"^ATM WITHDRAWAL", 1905, "auto_reconcile",
     "Existing model 21 covers 'ATM WITHDRAWAL ZBLMN\\d+' -> 1905 on journal 23. Widen model "
     "21's regex instead of adding a duplicate."),

    (70, "BLOM: ATM cash deposit", r"^ATM CASH DEPOSIT", 1905, "auto_reconcile", ""),

    (80, "BLOM: salaries", r"^SALAR(Y|IES)", 1559, "auto_reconcile",
     "4211 = account id 1559. Existing model 47 posts SALARY to 1703 (6310, expense) and must be "
     "repointed to 1559, or the two disagree. NOTE 1559 is typed asset_receivable, which is wrong "
     "for a 'due to personnel' account — flag to the accountant."),

    (90, "BLOM: POS card settlements", r"^SALES VOUCHERS", 1918, "auto_reconcile",
     "Clears the POS receivable raised when the card sale was booked."),

    (100, "BLOM: FX conversion", r"^FX OPR", 516, "manual",
     "7910 does NOT exist in company 2 (it is Presentail LTD account 202). The company-2 pair is "
     "516/675100 (loss) and 582/7751 (gain). Manual, because the other leg is the LBP journal."),

    (110, "BLOM: Alfa / MIC line bill", r"^MIC \d+ \(ALFA\) \d+", None, "none",
     "Already auto-matches vendor bill INV0051656588. Do not model it."),

    (120, "BLOM: incoming transfers", r"^IPO/", None, "none",
     "DO NOT auto-write-off. Existing model 42 sends every IPO/ to account 1927 'Loan Zakaria "
     "Ammache' at 100%. These lines are a MIX: shareholder funding from Adnan (Zakaria) Ammache, "
     "customer receipts (BLOOMS FLORA, partner 131) and a deposit refund (Zeal SAL, partner 238). "
     "Needs a payer-based split, never a blanket rule."),

    (125, "BLOM: online redemption credit", r"^CR FROM ONLINE REDEMPTION", None, "manual",
     "Card refunds. In April 2026 the two amounts (17.51, 8.99) exactly match TOTERS charges, so "
     "they look like reversals against 1661 — but that is an inference, not a fact. Keep manual "
     "until someone confirms what is being redeemed."),

    (130, "BLOM: Toters settlements", r"TOTERS", 1661, "auto_reconcile",
     "Models 20 and 22 share an IDENTICAL regex but post to 1661 vs 1885, both at sequence 10 — "
     "outcome undefined today. Take journal 23 off model 22 before enabling."),
]


def load_labels(path):
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        return [(r[2], r[3]) for r in ws.iter_rows(min_row=2, values_only=True)]
    import csv
    from decimal import Decimal
    out = []
    for r in csv.reader(open(path, encoding="utf-8-sig", newline="")):
        if len(r) > 4 and re.match(r"^\d{2}/\d{2}/\d{4}$", (r[0] or "").strip()):
            try:
                out.append((re.sub(r"\s+", " ", r[2].strip()).upper(),
                            Decimal((r[3] or "").replace(",", ""))))
            except Exception:
                pass
    return out


def classify(label):
    for seq, name, rx, acct, trig, _ in MODELS:
        if re.search(rx, label, re.I):
            return seq, name, acct, trig
    return None


def acct_str(a):
    if a is None:
        return "(invoice matching / human)"
    code, nm, _ = ACCOUNTS.get(a, ("?", "?", "?"))
    return f"{a} [{code}] {nm}"


def report(path):
    rows = load_labels(path)
    hits, misses = Counter(), []
    auto = manual = native = 0
    for label, amt in rows:
        m = classify(label)
        if not m:
            misses.append((label, amt)); continue
        hits[m[1]] += 1
        auto += m[3] == "auto_reconcile"; manual += m[3] == "manual"; native += m[3] == "none"
    print(f"\n{path}: {len(rows)} transactions")
    for seq, name, rx, acct, trig, _ in MODELS:
        if hits[name]:
            print(f"  [{seq:>3}] {name:<38} {hits[name]:>3}x -> {acct_str(acct)}  ({trig})")
    print(f"  auto {auto} | invoice matching {native} | human {manual} | UNCLASSIFIED {len(misses)}")
    for label, amt in misses:
        print(f"     ? {label[:60]:<60} {amt}")
    return len(misses)


def main():
    paths = sys.argv[1:] or ["jun.xlsx"]
    bad = sum(report(p) for p in paths)
    print("\n  Models already on journal 23 (collision check):")
    for mid, (nm, rx, acct) in EXISTING_ON_J23.items():
        print(f"    model {mid:>3} {nm:<28} {rx:<28} -> {acct}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
